"""Course files for an evaluation set: `data/<set>/corpus/<course>/<files>`.

Real sets hold the student's own PDF/DOCX/PPTX files. The seed set is kept as
text (`dbms-notes.pdf.txt`) and rendered into real files here, so evaluation runs through
the same extractors as uploads. Pages are separated by a line containing only `===`; the
first line of each page is its title (a DOCX Heading 1 or a slide title).
"""

import hashlib
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches, Pt

from .dataset import set_dir

BUILD_DIR = Path(__file__).parent / "build"
SUPPORTED = (".pdf", ".docx", ".pptx")


def course_files(set_name: str) -> dict[str, list[Path]]:
    """course -> its files, rendering text sources into real documents first."""
    courses: dict[str, list[Path]] = {}
    for course_dir in sorted(p for p in (set_dir(set_name) / "corpus").iterdir() if p.is_dir()):
        out = BUILD_DIR / set_name / course_dir.name
        out.mkdir(parents=True, exist_ok=True)
        files = [p for p in sorted(course_dir.iterdir()) if p.suffix.lower() in SUPPORTED]
        for text_file in sorted(course_dir.glob("*.txt")):
            target = out / text_file.stem  # "dbms-notes.pdf.txt" -> "dbms-notes.pdf"
            if not target.exists() or target.stat().st_mtime < text_file.stat().st_mtime:
                render(text_file.read_text(encoding="utf-8"), target)
            files.append(target)
        courses[course_dir.name] = files
    return courses


def corpus_fingerprint(set_name: str, course: str) -> str:
    """Changes whenever a course's source files change, so a stale index is never reused."""
    digest = hashlib.sha256()
    for path in sorted((set_dir(set_name) / "corpus" / course).iterdir()):
        digest.update(path.name.encode())
        digest.update(hashlib.sha256(path.read_bytes()).digest())
    return digest.hexdigest()[:12]


def render(text: str, target: Path) -> None:
    pages = [p.strip() for p in text.split("\n===\n") if p.strip()]
    {".pdf": _render_pdf, ".docx": _render_docx, ".pptx": _render_pptx}[target.suffix](
        pages, target
    )


def _render_pdf(pages: list[str], target: Path) -> None:
    document = pymupdf.open()
    for page_text in pages:
        page = document.new_page()  # A4
        title, _, body = page_text.partition("\n")
        page.insert_textbox(pymupdf.Rect(56, 50, 540, 80), title, fontsize=14)
        overflow = page.insert_textbox(pymupdf.Rect(56, 85, 540, 800), body, fontsize=9.5)
        if overflow < 0:
            raise ValueError(f"page text doesn't fit on one PDF page: {title!r}")
    document.save(target)
    document.close()


def _render_docx(pages: list[str], target: Path) -> None:
    document = DocxDocument()
    for page_text in pages:
        title, _, body = page_text.partition("\n")
        document.add_paragraph(title, style="Heading 1")
        for paragraph in body.split("\n\n"):
            document.add_paragraph(paragraph.strip())
    document.save(target)


def _render_pptx(pages: list[str], target: Path) -> None:
    presentation = Presentation()
    for page_text in pages:
        title, _, body = page_text.partition("\n")
        lines = body.split("\n")
        notes = [line.removeprefix("Notes:").strip() for line in lines if line.startswith("Notes:")]
        bullets = [line for line in lines if line.strip() and not line.startswith("Notes:")]
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])  # title only
        slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5.5))
        box.text_frame.word_wrap = True
        box.text_frame.text = "\n".join(bullets)
        for paragraph in box.text_frame.paragraphs:
            paragraph.font.size = Pt(14)
        if notes:
            slide.notes_slide.notes_text_frame.text = " ".join(notes)
    presentation.save(target)
