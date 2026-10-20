"""Text extraction from untrusted files. Runs only inside the worker (timeout + memory limit).

Each extractor returns pages in reading order:
- PDF: one entry per page
- PPTX: one entry per slide (including speaker notes and tables)
- DOCX: one entry per section (DOCX has no fixed pages; we split at Heading 1 or ~4000 chars)
"""

from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.models import DocumentKind

MAX_PAGES = 300
DOCX_SECTION_CHARS = 4000


class ExtractionError(Exception):
    """Expected failure with a message that is safe to show the user."""


@dataclass(frozen=True)
class Page:
    number: int
    text: str


def extract_pages(path: Path, kind: DocumentKind) -> list[Page]:
    extractor = {
        DocumentKind.PDF: _extract_pdf,
        DocumentKind.PPTX: _extract_pptx,
        DocumentKind.DOCX: _extract_docx,
    }[kind]
    pages = extractor(path)
    if len(pages) > MAX_PAGES:
        raise ExtractionError(f"Documents can have at most {MAX_PAGES} pages or slides.")
    return pages


def _extract_pdf(path: Path) -> list[Page]:
    try:
        document = pymupdf.open(path, filetype="pdf")
    except (pymupdf.FileDataError, RuntimeError):
        raise ExtractionError("This PDF is damaged and couldn't be opened.") from None
    with document:
        if document.needs_pass:
            raise ExtractionError("Password-protected PDFs aren't supported.")
        if document.page_count > MAX_PAGES:
            raise ExtractionError(f"Documents can have at most {MAX_PAGES} pages or slides.")
        return [
            Page(number=i, text=page.get_text("text", sort=True))
            for i, page in enumerate(document, start=1)
        ]


def _extract_pptx(path: Path) -> list[Page]:
    presentation = Presentation(str(path))
    pages = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            _collect_shape_text(shape, parts)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                parts.append(f"Notes: {notes.text}")
        pages.append(Page(number=number, text="\n".join(parts)))
    return pages


def _collect_shape_text(shape, parts: list[str]) -> None:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            _collect_shape_text(child, parts)
    elif shape.has_text_frame:
        if shape.text_frame.text.strip():
            parts.append(shape.text_frame.text)
    elif getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))


def _extract_docx(path: Path) -> list[Page]:
    document = DocxDocument(str(path))
    sections: list[list[str]] = [[]]
    length = 0
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text
            style = (block.style.name if block.style is not None else "") or ""
            starts_section = style in ("Title", "Heading 1")
        elif isinstance(block, Table):
            text = "\n".join(" | ".join(cell.text for cell in row.cells) for row in block.rows)
            starts_section = False
        else:
            continue
        if not text.strip():
            continue
        if sections[-1] and (starts_section or length + len(text) > DOCX_SECTION_CHARS):
            sections.append([])
            length = 0
        sections[-1].append(text)
        length += len(text)
    return [
        Page(number=i, text="\n".join(parts))
        for i, parts in enumerate((s for s in sections if s), start=1)
    ]
