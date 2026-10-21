"""Builders for real (and deliberately malicious) test documents, generated in memory."""

import io
import zipfile

import pymupdf
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches


def make_pdf(pages: list[str]) -> bytes:
    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        page.insert_text((72, 72), text, fontsize=11)
    data = document.tobytes()
    document.close()
    return data


def make_docx(paragraphs: list[tuple[str, str | None]]) -> bytes:
    """paragraphs: (text, style) pairs, e.g. ("Normalization", "Heading 1")."""
    document = DocxDocument()
    for text, style in paragraphs:
        document.add_paragraph(text, style=style)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pptx(slides: list[dict]) -> bytes:
    """slides: dicts with optional "title", "body", "table" (list of rows), "notes"."""
    presentation = Presentation()
    for spec in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])  # title only
        if "title" in spec:
            slide.shapes.title.text = spec["title"]
        if "body" in spec:
            box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
            box.text_frame.text = spec["body"]
        if "table" in spec:
            rows = spec["table"]
            shape = slide.shapes.add_table(
                len(rows), len(rows[0]), Inches(1), Inches(3), Inches(6), Inches(2)
            )
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    shape.table.cell(r, c).text = value
        if "notes" in spec:
            slide.notes_slide.notes_text_frame.text = spec["notes"]
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def make_zip(entries: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def make_docx_zip_bomb() -> bytes:
    """A structurally valid DOCX with one entry that inflates ~1000x (40 MB of zeros)."""
    return make_zip(
        {
            "[Content_Types].xml": b"<Types/>",
            "word/document.xml": b"<w:document/>",
            "word/media/bomb.bin": b"\0" * (40 * 1024 * 1024),
        }
    )


def make_docm_with_macro() -> bytes:
    return make_zip(
        {
            "[Content_Types].xml": b"<Types/>",
            "word/document.xml": b"<w:document/>",
            "word/vbaProject.bin": b'Attribute VB_Name = "evil"',
        }
    )


def make_scanned_pdf(pages: list[str | None]) -> bytes:
    """Each entry is typed text, or None for a page that is only a picture (a scan)."""
    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        if text is None:
            picture = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 80), 0)
            picture.clear_with(180)
            page.insert_image(pymupdf.Rect(72, 72, 360, 456), pixmap=picture)
        else:
            page.insert_text((72, 72), text, fontsize=11)
    data = document.tobytes()
    document.close()
    return data


def make_photo(fmt: str = "JPEG", size: tuple[int, int] = (300, 200), rotate=False) -> bytes:
    """A photo as a phone saves it; `rotate` stores it sideways with an EXIF orientation."""
    from PIL import Image

    photo = Image.new("RGB", size, (240, 236, 225))
    out = io.BytesIO()
    if rotate:
        exif = Image.Exif()
        exif[0x0112] = 6  # Orientation: rotate 90° clockwise to display
        photo.save(out, fmt, exif=exif)
    else:
        photo.save(out, fmt)
    return out.getvalue()


def make_math_pdf(math_pages: set[int], pages: int = 3) -> bytes:
    """Typed pages; those in `math_pages` also carry equations set in the Symbol font."""
    document = pymupdf.open()
    for number in range(1, pages + 1):
        page = document.new_page()
        page.insert_text((72, 72), f"Lecture notes, page {number}: kinetic energy.", fontsize=11)
        if number in math_pages:
            for i in range(6):  # Greek and operators in a math font, as TeX output has
                page.insert_text((72, 110 + 24 * i), "a b g d S p", fontname="symb", fontsize=12)
    data = document.tobytes()
    document.close()
    return data
